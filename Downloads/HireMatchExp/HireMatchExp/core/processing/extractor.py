
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
import io
from pathlib import Path
from typing import Any, BinaryIO, Dict, List, Optional, Union
import xml.etree.ElementTree as ET
import zipfile

from core.models.enums import FileType
from core.processing.cleaner import TextCleaner
from core.utils.errors import ExtractionError
from core.utils.logger import logger


@dataclass
class PageText:
    page_number: int
    raw_text: str
    cleaned_text: str
    char_count: int


@dataclass
class ExtractedDocument:
    raw_text: str
    cleaned_text: str
    pages: List[PageText] = field(default_factory=list)
    tables_text: List[str] = field(default_factory=list)
    is_empty: bool = False
    ocr_needed: bool = False
    metadata: Dict[str, Any] = field(default_factory=dict)

    @property
    def text(self) -> str:
        return self.cleaned_text


class BaseExtractor(ABC):

    @abstractmethod
    def extract_document(self, file_source: Union[bytes, BinaryIO, Path, str]) -> ExtractedDocument:
        pass

    def extract_text(self, file_source: Union[bytes, BinaryIO, Path, str]) -> str:
        doc = self.extract_document(file_source)
        return doc.cleaned_text

    def extract(self, file_source: Union[bytes, BinaryIO, Path, str], filename: Optional[str] = None) -> ExtractedDocument:
        return self.extract_document(file_source)


class PyMuPDFExtractor(BaseExtractor):

    def extract_document(self, file_source: Union[bytes, BinaryIO, Path, str]) -> ExtractedDocument:
        try:
            import fitz
        except ImportError:
            logger.warning("PyMuPDF (fitz) is not installed. Using raw byte parser fallback.")
            return self._fallback_pdf_extraction(file_source)

        try:
            doc = None
            if isinstance(file_source, bytes):
                doc = fitz.open(stream=file_source, filetype="pdf")
            elif isinstance(file_source, (str, Path)):
                doc = fitz.open(str(file_source))
            elif hasattr(file_source, "read"):
                data = file_source.read()
                doc = fitz.open(stream=data, filetype="pdf")
            else:
                raise ExtractionError(f"Unsupported file source type: {type(file_source)}")

            page_records: List[PageText] = []
            full_raw_parts: List[str] = []
            total_chars = 0

            for page_idx, page in enumerate(doc):
                page_raw = page.get_text() or ""
                page_clean = TextCleaner.clean(page_raw)
                char_count = len(page_raw)
                total_chars += char_count

                full_raw_parts.append(page_raw)
                page_records.append(
                    PageText(
                        page_number=page_idx + 1,
                        raw_text=page_raw,
                        cleaned_text=page_clean,
                        char_count=char_count,
                    )
                )

            page_count = len(doc)
            meta = {
                "format": "PDF",
                "page_count": page_count,
                "author": doc.metadata.get("author") if doc.metadata else None,
                "title": doc.metadata.get("title") if doc.metadata else None,
            }
            doc.close()

            full_raw = "\n\n".join(full_raw_parts).strip()
            full_clean = TextCleaner.clean(full_raw)

            if page_count == 0 and not self._is_empty_source(file_source):
                logger.info(f"PyMuPDF found 0 pages. Attempting raw fallback extraction.")
                fallback_doc = self._fallback_pdf_extraction(file_source)
                if fallback_doc.cleaned_text:
                    return fallback_doc

            is_empty = total_chars == 0 or len(full_clean) == 0
            ocr_needed = False
            if page_count > 0 and (total_chars < 30 or len(full_clean) < 30):
                ocr_needed = True
                logger.info(f"PDF has {page_count} pages but only {total_chars} characters. OCR is required.")

            return ExtractedDocument(
                raw_text=full_raw,
                cleaned_text=full_clean,
                pages=page_records,
                is_empty=is_empty,
                ocr_needed=ocr_needed,
                metadata=meta,
            )

        except Exception as e:
            logger.warning(f"PyMuPDF extraction failed: {e}. Falling back to byte extraction.")
            try:
                return self._fallback_pdf_extraction(file_source)
            except Exception:
                raise ExtractionError(f"Failed to extract text from PDF: {e}") from e

    def _is_empty_source(self, file_source: Union[bytes, BinaryIO, Path, str]) -> bool:
        try:
            if isinstance(file_source, bytes):
                return len(file_source) == 0
            elif isinstance(file_source, (str, Path)):
                return Path(file_source).stat().st_size == 0
            elif hasattr(file_source, "seek") and hasattr(file_source, "tell"):
                pos = file_source.tell()
                file_source.seek(0, io.SEEK_END)
                size = file_source.tell()
                file_source.seek(pos)
                return size == 0
        except Exception:
            pass
        return False

    def _fallback_pdf_extraction(self, file_source: Union[bytes, BinaryIO, Path, str]) -> ExtractedDocument:
        data = b""
        if isinstance(file_source, bytes):
            data = file_source
        elif isinstance(file_source, (str, Path)):
            with open(file_source, "rb") as f:
                data = f.read()
        elif hasattr(file_source, "read"):
            data = file_source.read()
            if hasattr(file_source, "seek"):
                file_source.seek(0)

        if b"/Filter" in data and (b"/DCTDecode" in data or b"/FlateDecode" in data):
            logger.info("Fallback byte extraction encountered compressed binary PDF stream. Flagging OCR needed.")
            return ExtractedDocument(
                raw_text="",
                cleaned_text="",
                pages=[],
                is_empty=True,
                ocr_needed=True,
                metadata={"format": "PDF", "fallback_used": True, "ocr_required": True},
            )

        text = "".join(chr(b) for b in data if 32 <= b <= 126 or b in (10, 13))
        cleaned = TextCleaner.clean(text)

        is_empty = len(cleaned) == 0
        ocr_needed = len(cleaned) < 30

        return ExtractedDocument(
            raw_text=text,
            cleaned_text=cleaned,
            pages=[PageText(page_number=1, raw_text=text, cleaned_text=cleaned, char_count=len(text))],
            is_empty=is_empty,
            ocr_needed=ocr_needed,
            metadata={"format": "PDF", "fallback_used": True},
        )


class DocxExtractor(BaseExtractor):

    def extract_document(self, file_source: Union[bytes, BinaryIO, Path, str]) -> ExtractedDocument:
        try:
            import docx
            file_stream = self._to_stream(file_source)
            doc = docx.Document(file_stream)

            paragraphs: List[str] = [p.text for p in doc.paragraphs if p.text.strip()]
            tables_data: List[str] = []

            for t_idx, table in enumerate(doc.tables):
                table_lines = []
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        table_lines.append(" | ".join(cells))
                if table_lines:
                    tables_data.append("\n".join(table_lines))

            all_parts = list(paragraphs)
            if tables_data:
                all_parts.append("\n--- Table Data ---\n" + "\n\n".join(tables_data))

            full_raw = "\n\n".join(all_parts).strip()
            full_clean = TextCleaner.clean(full_raw)

            return ExtractedDocument(
                raw_text=full_raw,
                cleaned_text=full_clean,
                pages=[PageText(page_number=1, raw_text=full_raw, cleaned_text=full_clean, char_count=len(full_raw))],
                tables_text=tables_data,
                is_empty=(len(full_clean) == 0),
                ocr_needed=False,
                metadata={"format": "DOCX", "paragraphs_count": len(paragraphs), "tables_count": len(tables_data)},
            )
        except ImportError:
            logger.info("python-docx not installed. Falling back to native XML extraction.")
            return self._extract_via_zip_xml(file_source)
        except Exception as e:
            logger.warning(f"python-docx parsing error: {e}. Trying XML fallback.")
            return self._extract_via_zip_xml(file_source)

    def _to_stream(self, file_source: Union[bytes, BinaryIO, Path, str]) -> BinaryIO:
        if isinstance(file_source, bytes):
            return io.BytesIO(file_source)
        elif isinstance(file_source, (str, Path)):
            with open(file_source, "rb") as f:
                return io.BytesIO(f.read())
        elif hasattr(file_source, "read"):
            data = file_source.read()
            return io.BytesIO(data)
        raise ExtractionError(f"Unsupported source type: {type(file_source)}")

    def _extract_via_zip_xml(self, file_source: Union[bytes, BinaryIO, Path, str]) -> ExtractedDocument:
        try:
            stream = self._to_stream(file_source)
            with zipfile.ZipFile(stream) as z:
                xml_content = z.read("word/document.xml")

            tree = ET.fromstring(xml_content)
            texts = []
            for elem in tree.iter():
                if elem.tag.endswith("}t") and elem.text:
                    texts.append(elem.text)
                elif elem.tag.endswith("}p"):
                    texts.append("\n")

            full_raw = "".join(texts).strip()
            full_clean = TextCleaner.clean(full_raw)

            return ExtractedDocument(
                raw_text=full_raw,
                cleaned_text=full_clean,
                pages=[PageText(page_number=1, raw_text=full_raw, cleaned_text=full_clean, char_count=len(full_raw))],
                is_empty=(len(full_clean) == 0),
                ocr_needed=False,
                metadata={"format": "DOCX", "fallback_used": True},
            )
        except Exception as e:
            logger.error(f"DOCX XML extraction failed: {e}")
            raise ExtractionError(f"Failed to extract text from DOCX: {e}") from e


class DocumentExtractorFactory:

    @staticmethod
    def get_extractor(file_type: FileType) -> BaseExtractor:
        if file_type == FileType.PDF:
            return PyMuPDFExtractor()
        elif file_type == FileType.DOCX:
            return DocxExtractor()
        raise ExtractionError(f"No extractor registered for file type: {file_type}")
