
import io
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd

from core.models.report_models import (
    CandidateComparisonData,
    CandidateReportData,
    SessionSummaryReportData,
    ShortlistReportData,
)


class XLSXReportFormatter:

    HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="111827")
    REGULAR_FONT = Font(name="Calibri", size=10, color="1F2937")
    MUTED_FONT = Font(name="Calibri", size=9, color="6B7280", italic=True)

    @classmethod
    def format_candidate_report(cls, data: CandidateReportData) -> bytes:
        wb = openpyxl.Workbook()

        ws_summary = wb.active
        ws_summary.title = "Executive Summary"
        cls._apply_candidate_summary_sheet(ws_summary, data)

        ws_crit = wb.create_sheet(title="Criterion Breakdown")
        cls._apply_candidate_criteria_sheet(ws_crit, data)

        ws_bg = wb.create_sheet(title="Extracted Profile")
        cls._apply_candidate_background_sheet(ws_bg, data)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @classmethod
    def format_session_report(cls, data: SessionSummaryReportData) -> bytes:
        wb = openpyxl.Workbook()

        ws_metrics = wb.active
        ws_metrics.title = "Session Overview"
        cls._apply_session_metrics_sheet(ws_metrics, data)

        ws_cands = wb.create_sheet(title="Candidate Manifest")
        cls._apply_session_candidates_sheet(ws_cands, data)

        ws_crits = wb.create_sheet(title="Criterion Pool Analytics")
        cls._apply_session_criteria_sheet(ws_crits, data)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @classmethod
    def format_comparison_report(cls, data: CandidateComparisonData) -> bytes:
        wb = openpyxl.Workbook()

        ws_comp = wb.active
        ws_comp.title = "Comparison Matrix"
        cls._apply_comparison_grid_sheet(ws_comp, data)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @classmethod
    def format_shortlist_report(cls, data: ShortlistReportData) -> bytes:
        wb = openpyxl.Workbook()

        ws_sl = wb.active
        ws_sl.title = "Shortlist Review Queue"
        cls._apply_shortlist_sheet(ws_sl, data)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @classmethod
    def _apply_candidate_summary_sheet(cls, ws: Any, data: CandidateReportData):
        ws.cell(row=1, column=1, value=f"Candidate Evaluation Dossier: {data.candidate_reference}").font = cls.TITLE_FONT
        ws.cell(row=2, column=1, value=f"Target Job: {data.job_title} | Session: {data.session_name}").font = cls.MUTED_FONT

        rows = [
            ("Original Filename", data.original_filename),
            ("Overall Score", f"{data.overall_score:.1f}%" if data.overall_score is not None else "N/A"),
            ("Evidence Coverage", f"{data.evidence_coverage:.1f}%"),
            ("Completeness Status", data.completeness_status),
            ("Scoring Config Version", data.scoring_config_version),
            ("Recruiter Review Status", data.review_status),
            ("Recruiter Notes", data.reviewer_notes or "None"),
            ("Supported Criteria", f"{data.supported_criteria_count} of {data.total_criteria_count}"),
            ("Conflicting Criteria", str(data.conflicting_criteria_count)),
            ("Resume Quality Score", f"{data.resume_quality_score:.1f}%" if data.resume_quality_score is not None else "N/A"),
            ("Resume Quality Label", data.resume_quality_label or "N/A"),
            ("Report Generated At", data.generated_at),
        ]
        for idx, (label, val) in enumerate(rows, start=4):
            ws.cell(row=idx, column=1, value=label).font = Font(name="Calibri", bold=True, size=10)
            ws.cell(row=idx, column=2, value=val).font = cls.REGULAR_FONT

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 45

    @classmethod
    def _apply_candidate_criteria_sheet(cls, ws: Any, data: CandidateReportData):
        headers = ["Criterion Name", "Category", "Type", "Weight", "Support Value", "Normalized Points", "Match Status", "Explanation"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT

        for r_idx, c in enumerate(data.criterion_items, start=2):
            cat_name = c.category.value if hasattr(c.category, "value") else str(c.category)
            t_name = c.criterion_type.value if hasattr(c.criterion_type, "value") else str(c.criterion_type)
            st_name = c.match_status.value if hasattr(c.match_status, "value") else str(c.match_status)

            ws.cell(row=r_idx, column=1, value=c.criterion_name).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=2, value=cat_name).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=3, value=t_name).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=4, value=c.configured_weight).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=5, value=c.support_value).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=6, value=c.normalized_points).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=7, value=st_name).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=8, value=c.explanation).font = cls.REGULAR_FONT

        for col in ["A", "B", "C", "D", "E", "F", "G", "H"]:
            ws.column_dimensions[col].width = 20
        ws.column_dimensions["H"].width = 40

    @classmethod
    def _apply_candidate_background_sheet(cls, ws: Any, data: CandidateReportData):
        ws.cell(row=1, column=1, value="Extracted Skills").font = cls.TITLE_FONT
        for idx, s in enumerate(data.extracted_skills, start=2):
            ws.cell(row=idx, column=1, value=s).font = cls.REGULAR_FONT

        next_row = len(data.extracted_skills) + 4
        ws.cell(row=next_row, column=1, value="Work Experience").font = cls.TITLE_FONT
        for idx, exp in enumerate(data.work_experience, start=next_row + 1):
            title = exp.get("job_title") or exp.get("title", "Role")
            company = exp.get("employer") or exp.get("company", "Company")
            ws.cell(row=idx, column=1, value=f"{title} - {company}").font = cls.REGULAR_FONT

    @classmethod
    def _apply_session_metrics_sheet(cls, ws: Any, data: SessionSummaryReportData):
        ws.cell(row=1, column=1, value=f"Screening Session Summary: {data.session_name}").font = cls.TITLE_FONT
        ws.cell(row=2, column=1, value=f"Job: {data.job_title} ({data.department}) | Resumes: {data.total_resumes_enrolled}").font = cls.MUTED_FONT

        rows = [
            ("Total Resumes Enrolled", data.total_resumes_enrolled),
            ("Completed Processing", data.completed_processing_count),
            ("Failed Processing", data.failed_processing_count),
            ("Completed Evaluations", data.completed_evaluation_count),
            ("Mean Score", f"{data.score_distribution.get('mean', 0.0):.1f}%" if data.score_distribution.get('mean') is not None else "N/A"),
            ("Median Score", f"{data.score_distribution.get('median', 0.0):.1f}%" if data.score_distribution.get('median') is not None else "N/A"),
            ("Shortlisted Count", data.shortlist_summary.get("SHORTLISTED", 0)),
            ("In Review Count", data.shortlist_summary.get("IN_REVIEW", 0)),
        ]
        for idx, (lbl, val) in enumerate(rows, start=4):
            ws.cell(row=idx, column=1, value=lbl).font = Font(name="Calibri", bold=True, size=10)
            ws.cell(row=idx, column=2, value=val).font = cls.REGULAR_FONT

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 30

    @classmethod
    def _apply_session_candidates_sheet(cls, ws: Any, data: SessionSummaryReportData):
        headers = ["Candidate Ref", "Original Filename", "Processing Status", "Overall Score", "Resume Quality", "Quality Label", "Coverage %", "Review Status", "Reviewer Notes"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT

        for r_idx, c in enumerate(data.candidate_summaries, start=2):
            q_score_str = f"{c['resume_quality_score']:.1f}%" if c.get("resume_quality_score") is not None else ""
            q_label_str = c.get("resume_quality_label") or ""

            ws.cell(row=r_idx, column=1, value=c["candidate_reference"]).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=2, value=c["original_filename"]).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=3, value=c["processing_status"]).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=4, value=c["overall_score"] if c["overall_score"] is not None else "").font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=5, value=q_score_str).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=6, value=q_label_str).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=7, value=c["evidence_coverage"]).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=8, value=c["review_status"]).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=9, value=c["reviewer_notes"]).font = cls.REGULAR_FONT

        for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I"]:
            ws.column_dimensions[col].width = 22

    @classmethod
    def _apply_session_criteria_sheet(cls, ws: Any, data: SessionSummaryReportData):
        headers = ["Criterion Name", "Category", "Type", "Supported %", "Partial %", "Insufficient %", "Conflicting %"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT

        for r_idx, cp in enumerate(data.criterion_performance, start=2):
            ws.cell(row=r_idx, column=1, value=cp["criterion_name"]).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=2, value=cp["category"]).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=3, value=cp["criterion_type"]).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=4, value=cp["supported_pct"]).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=5, value=cp["partial_pct"]).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=6, value=cp["insufficient_pct"]).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=7, value=cp["conflicting_pct"]).font = cls.REGULAR_FONT

        for col in ["A", "B", "C", "D", "E", "F", "G"]:
            ws.column_dimensions[col].width = 20

    @classmethod
    def _apply_comparison_grid_sheet(cls, ws: Any, data: CandidateComparisonData):
        ws.cell(row=1, column=1, value=f"Candidate Comparison Matrix: {data.job_title}").font = cls.TITLE_FONT

        headers = ["Criterion Name", "Weight"]
        for h in data.candidate_headers:
            headers.extend([f"{h['candidate_reference']} Status", f"{h['candidate_reference']} Points"])

        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col_idx, value=h_text)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT

        for r_idx, row in enumerate(data.aligned_criteria_rows, start=4):
            ws.cell(row=r_idx, column=1, value=row.criterion_name).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=2, value=row.configured_weight).font = cls.REGULAR_FONT
            col_ptr = 3
            for h in data.candidate_headers:
                eval_item = row.candidate_evaluations.get(h["candidate_id"], {})
                ws.cell(row=r_idx, column=col_ptr, value=eval_item.get("match_status", "N/A")).font = cls.REGULAR_FONT
                ws.cell(row=r_idx, column=col_ptr + 1, value=eval_item.get("normalized_points", 0.0)).font = cls.REGULAR_FONT
                col_ptr += 2

        for col_idx in range(1, len(headers) + 1):
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 20

    @classmethod
    def _apply_shortlist_sheet(cls, ws: Any, data: ShortlistReportData):
        headers = ["Candidate Ref", "Filename", "Review Status", "Overall Score", "Coverage %", "Key Highlights", "Reviewer Notes", "Last Updated"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT

        for r_idx, item in enumerate(data.items, start=2):
            st_val = item.review_status.value if hasattr(item.review_status, "value") else str(item.review_status)
            ws.cell(row=r_idx, column=1, value=item.candidate_reference).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=2, value=item.original_filename).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=3, value=st_val).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=4, value=item.overall_score if item.overall_score is not None else "").font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=5, value=item.evidence_coverage).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=6, value=" | ".join(item.key_evidence_highlights)).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=7, value=item.reviewer_notes).font = cls.REGULAR_FONT
            ws.cell(row=r_idx, column=8, value=item.updated_at).font = cls.REGULAR_FONT

        for col in ["A", "B", "C", "D", "E", "F", "G", "H"]:
            ws.column_dimensions[col].width = 22
