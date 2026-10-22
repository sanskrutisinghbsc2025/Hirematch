
import csv
from datetime import datetime, timezone
import io
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import uuid

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import AppConfig, settings
from core.analytics.models import (
    CriterionAnalytics,
    PortfolioAnalytics,
    SessionAnalytics,
    SessionComparisonAnalytics,
    SkillCoverageItem,
    TraceableInsight,
)
from core.models.report_models import ExportFormat
from core.utils.logger import logger


class AnalyticsExportService:

    METRIC_DEFINITIONS = {
        "Overall Mean Score": "Arithmetic average of valid composite candidate scores (0-100 scale).",
        "Overall Median Score": "50th percentile of candidate scores; unaffected by extreme score outliers.",
        "Score Interquartile Range (IQR)": "Spread between Q3 (75th percentile) and Q1 (25th percentile); measures score dispersion.",
        "Evidence Coverage Rate (%)": "Percentage of evaluated criteria where candidate provided verified SUPPORTED or PARTIAL evidence.",
        "Insufficient Rate (%)": "Percentage of candidates where requirement was NOT_FOUND or INSUFFICIENT_EVIDENCE.",
        "Criterion Variance (\u03c3\u00b2)": "Degree of spread in candidate support values for a requirement. Low variance (<0.006) indicates minimal differentiation.",
        "Keyword vs Semantic Divergence": "Mean absolute score difference between exact keyword search and dense neural embeddings.",
        "Acute Skill Gap": "Job requirement where >60% of candidates in the pool lack verified evidence.",
        "Shortlist Conversion Rate (%)": "Percentage of total enrolled candidates moved to SHORTLISTED review status by a recruiter.",
    }

    def __init__(self, config: Optional[AppConfig] = None):
        self.config = config or settings
        self.export_dir = self.config.reports_dir / "analytics"
        self.export_dir.mkdir(parents=True, exist_ok=True)

    def export_session_analytics_csv(
        self, session_analytics: SessionAnalytics, filters_meta: Optional[Dict[str, Any]] = None
    ) -> Tuple[bytes, str]:
        buf = io.StringIO()
        writer = csv.writer(buf, lineterminator="\n")

        writer.writerow(["# HIREMATCH BUSINESS ANALYTICS EXPORT"])
        writer.writerow(["# EXPORT TYPE", "Screening Session Analytics"])
        writer.writerow(["# SESSION ID", session_analytics.session_id])
        writer.writerow(["# SESSION NAME", session_analytics.session_name])
        writer.writerow(["# JOB REQUISITION", session_analytics.job_title])
        writer.writerow(["# GENERATED AT", session_analytics.generated_at])
        writer.writerow(["# SCORING CONFIG VERSION", session_analytics.scoring_config_version])
        if filters_meta:
            for k, v in filters_meta.items():
                if v is not None:
                    writer.writerow([f"# FILTER: {k.upper()}", str(v)])
        writer.writerow([])

        writer.writerow(["# METRIC DEFINITIONS"])
        for k, v in self.METRIC_DEFINITIONS.items():
            writer.writerow(["#", k, v])
        writer.writerow([])

        writer.writerow(["--- SESSION OVERVIEW KPIS ---"])
        writer.writerow(["Metric", "Value"])
        writer.writerow(["Total Resumes Enrolled", session_analytics.total_enrolled])
        writer.writerow(["Total Scored Candidates", session_analytics.total_scored])
        writer.writerow(["Unscorable Count", session_analytics.unscorable_count])
        writer.writerow(["Mean Score (%)", session_analytics.score_mean if session_analytics.score_mean is not None else "N/A"])
        writer.writerow(["Median Score (%)", session_analytics.score_median if session_analytics.score_median is not None else "N/A"])
        writer.writerow(["Score Standard Deviation (\u03c3)", session_analytics.score_std_dev])
        writer.writerow(["Score IQR", session_analytics.score_iqr])
        writer.writerow(["Mean Evidence Coverage (%)", session_analytics.evidence_coverage_mean])
        writer.writerow([])

        writer.writerow(["--- CRITERION PERFORMANCE & VARIANCE ---"])
        writer.writerow([
            "criterion_id",
            "criterion_name",
            "category",
            "type",
            "weight",
            "mean_support_value",
            "mean_points",
            "std_dev",
            "variance",
            "is_low_variation",
            "evidence_coverage_rate_pct",
            "insufficient_rate_pct",
            "keyword_semantic_divergence",
            "semantic_boost_count",
        ])
        for ca in session_analytics.criterion_analytics:
            writer.writerow([
                ca.criterion_id,
                ca.criterion_name,
                ca.category,
                ca.criterion_type,
                ca.configured_weight,
                ca.mean_support_value,
                ca.mean_normalized_points,
                ca.std_dev,
                ca.variance,
                "YES" if ca.is_low_variation else "NO",
                ca.evidence_coverage_rate,
                ca.insufficient_rate,
                ca.keyword_vs_semantic_divergence,
                ca.semantic_boost_count,
            ])
        writer.writerow([])

        writer.writerow(["--- SKILL COVERAGE & GAPS ---"])
        writer.writerow([
            "canonical_name",
            "raw_skill_name",
            "category",
            "requirement_type",
            "evidence_count",
            "total_candidates",
            "pool_coverage_pct",
            "gap_status",
        ])
        for sc in session_analytics.skill_coverage:
            writer.writerow([
                sc.canonical_name,
                sc.raw_skill_name,
                sc.category,
                sc.requirement_type,
                sc.candidate_evidence_count,
                sc.total_candidates,
                sc.pool_coverage_percentage,
                sc.gap_status.value if hasattr(sc.gap_status, "value") else str(sc.gap_status),
            ])

        csv_bytes = buf.getvalue().encode("utf-8")
        filename = f"analytics_session_{session_analytics.session_id[:8]}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        file_path = str(self.export_dir / filename)
        with open(file_path, "wb") as f:
            f.write(csv_bytes)

        return csv_bytes, file_path

    def export_session_analytics_xlsx(
        self, session_analytics: SessionAnalytics, filters_meta: Optional[Dict[str, Any]] = None
    ) -> Tuple[bytes, str]:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        accent_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        font_title = Font(name="Arial", size=14, bold=True, color="0F172A")
        font_regular = Font(name="Arial", size=10)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        ws1 = wb.create_sheet(title="Executive Overview")
        ws1["A1"] = f"HireMatch Analytics: {session_analytics.session_name}"
        ws1["A1"].font = font_title
        ws1["A2"] = f"Position: {session_analytics.job_title} | Session ID: {session_analytics.session_id}"
        ws1["A2"].font = font_regular

        row = 4
        ws1.cell(row=row, column=1, value="Recruitment Performance Metric").font = font_header
        ws1.cell(row=row, column=1).fill = accent_fill
        ws1.cell(row=row, column=2, value="Aggregated Value").font = font_header
        ws1.cell(row=row, column=2).fill = accent_fill

        kpis = [
            ("Total Resumes Enrolled", session_analytics.total_enrolled),
            ("Total Candidates Scored", session_analytics.total_scored),
            ("Unscorable Resumes", session_analytics.unscorable_count),
            ("Mean Composite Score (%)", f"{session_analytics.score_mean:.1f}%" if session_analytics.score_mean is not None else "N/A"),
            ("Median Composite Score (%)", f"{session_analytics.score_median:.1f}%" if session_analytics.score_median is not None else "N/A"),
            ("Score Standard Deviation (\u03c3)", session_analytics.score_std_dev),
            ("Score Interquartile Range (IQR)", session_analytics.score_iqr),
            ("Tightly Clustered Alert", "YES" if session_analytics.is_tightly_clustered else "NO"),
            ("Mean Evidence Coverage Rate (%)", f"{session_analytics.evidence_coverage_mean:.1f}%"),
            ("Shortlisted Candidates", session_analytics.shortlist_counts.get("SHORTLISTED", 0)),
            ("In Review Candidates", session_analytics.shortlist_counts.get("IN_REVIEW", 0)),
            ("Scoring Config Version", session_analytics.scoring_config_version),
            ("Export Timestamp", session_analytics.generated_at),
        ]
        for name, val in kpis:
            row += 1
            ws1.cell(row=row, column=1, value=name).font = font_regular
            ws1.cell(row=row, column=2, value=val).font = font_regular
            ws1.cell(row=row, column=1).border = thin_border
            ws1.cell(row=row, column=2).border = thin_border

        ws2 = wb.create_sheet(title="Criterion Variance & AI")
        c_headers = [
            "Criterion Name",
            "Category",
            "Type",
            "Weight",
            "Mean Support",
            "Mean Points",
            "Std Dev (\u03c3)",
            "Variance",
            "Low Variation?",
            "Evidence Coverage %",
            "Insufficient %",
            "AI Divergence",
            "Semantic Boosts",
        ]
        for col_idx, h in enumerate(c_headers, start=1):
            cell = ws2.cell(row=1, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill

        r_idx = 2
        for ca in session_analytics.criterion_analytics:
            vals = [
                ca.criterion_name,
                ca.category,
                ca.criterion_type,
                ca.configured_weight,
                ca.mean_support_value,
                ca.mean_normalized_points,
                ca.std_dev,
                ca.variance,
                "YES" if ca.is_low_variation else "NO",
                ca.evidence_coverage_rate,
                ca.insufficient_rate,
                ca.keyword_vs_semantic_divergence,
                ca.semantic_boost_count,
            ]
            for col_idx, v in enumerate(vals, start=1):
                c = ws2.cell(row=r_idx, column=col_idx, value=v)
                c.font = font_regular
                c.border = thin_border
                if r_idx % 2 == 1:
                    c.fill = alt_fill
            r_idx += 1

        ws3 = wb.create_sheet(title="Skill Coverage & Gaps")
        s_headers = [
            "Canonical Skill",
            "Raw Skill in JD",
            "Category",
            "Requirement Type",
            "Evidence Count",
            "Total Candidates",
            "Pool Coverage %",
            "Gap Status",
        ]
        for col_idx, h in enumerate(s_headers, start=1):
            cell = ws3.cell(row=1, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill

        r_idx = 2
        for sc in session_analytics.skill_coverage:
            vals = [
                sc.canonical_name,
                sc.raw_skill_name,
                sc.category,
                sc.requirement_type,
                sc.candidate_evidence_count,
                sc.total_candidates,
                sc.pool_coverage_percentage,
                sc.gap_status.value if hasattr(sc.gap_status, "value") else str(sc.gap_status),
            ]
            for col_idx, v in enumerate(vals, start=1):
                c = ws3.cell(row=r_idx, column=col_idx, value=v)
                c.font = font_regular
                c.border = thin_border
                if r_idx % 2 == 1:
                    c.fill = alt_fill
            r_idx += 1

        ws4 = wb.create_sheet(title="Screening Insights")
        i_headers = ["Category", "Severity", "Insight Title", "Description", "Calculation Rule", "Suggested Action"]
        for col_idx, h in enumerate(i_headers, start=1):
            cell = ws4.cell(row=1, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill

        r_idx = 2
        for ins in session_analytics.insights:
            vals = [
                ins.category.value if hasattr(ins.category, "value") else str(ins.category),
                ins.severity.value if hasattr(ins.severity, "value") else str(ins.severity),
                ins.title,
                ins.description,
                ins.calculation_basis,
                ins.suggested_recruiter_action,
            ]
            for col_idx, v in enumerate(vals, start=1):
                c = ws4.cell(row=r_idx, column=col_idx, value=v)
                c.font = font_regular
                c.border = thin_border
                if r_idx % 2 == 1:
                    c.fill = alt_fill
            r_idx += 1

        ws5 = wb.create_sheet(title="Metric Definitions")
        ws5.cell(row=1, column=1, value="Metric Term").font = font_header
        ws5.cell(row=1, column=1).fill = accent_fill
        ws5.cell(row=1, column=2, value="Formal Mathematical Definition & Recruitment Purpose").font = font_header
        ws5.cell(row=1, column=2).fill = accent_fill

        r_idx = 2
        for term, definition in self.METRIC_DEFINITIONS.items():
            ws5.cell(row=r_idx, column=1, value=term).font = font_regular
            ws5.cell(row=r_idx, column=2, value=definition).font = font_regular
            ws5.cell(row=r_idx, column=1).border = thin_border
            ws5.cell(row=r_idx, column=2).border = thin_border
            r_idx += 1

        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

        bio = io.BytesIO()
        wb.save(bio)
        xlsx_bytes = bio.getvalue()

        filename = f"analytics_session_{session_analytics.session_id[:8]}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path = str(self.export_dir / filename)
        with open(file_path, "wb") as f:
            f.write(xlsx_bytes)

        return xlsx_bytes, file_path

    def export_portfolio_analytics_csv(
        self, portfolio_analytics: PortfolioAnalytics, filters_meta: Optional[Dict[str, Any]] = None
    ) -> Tuple[bytes, str]:
        buf = io.StringIO()
        writer = csv.writer(buf, lineterminator="\n")

        writer.writerow(["# HIREMATCH PORTFOLIO ANALYTICS EXPORT"])
        writer.writerow(["# GENERATED AT", portfolio_analytics.generated_at])
        writer.writerow(["# TOTAL SCREENING SESSIONS", portfolio_analytics.total_sessions])
        writer.writerow(["# ACTIVE JOBS COUNT", portfolio_analytics.active_jobs_count])
        writer.writerow([])

        writer.writerow(["--- PORTFOLIO KPIS ---"])
        writer.writerow(["Metric", "Value"])
        writer.writerow(["Total Resumes Enrolled", portfolio_analytics.total_resumes_enrolled])
        writer.writerow(["Total Resumes Processed", portfolio_analytics.total_resumes_processed])
        writer.writerow(["Completed Extraction", portfolio_analytics.completed_processing_count])
        writer.writerow(["Partial Extraction", portfolio_analytics.partial_processing_count])
        writer.writerow(["Failed Extraction", portfolio_analytics.failed_processing_count])
        writer.writerow(["Duplicate Resumes", portfolio_analytics.duplicate_processing_count])
        writer.writerow(["Avg Session Processing Seconds", portfolio_analytics.average_session_processing_seconds or "N/A"])
        writer.writerow(["Overall Mean Score (%)", portfolio_analytics.overall_mean_score or "N/A"])
        writer.writerow(["Overall Median Score (%)", portfolio_analytics.overall_median_score or "N/A"])
        writer.writerow(["Score Quartile Q1 (25%)", portfolio_analytics.score_quartiles.get("q1", 0.0)])
        writer.writerow(["Score Quartile Q3 (75%)", portfolio_analytics.score_quartiles.get("q3", 0.0)])
        writer.writerow([])

        writer.writerow(["--- TIME SERIES SCREENING TRENDS ---"])
        writer.writerow(["period", "session_count", "resume_count", "avg_score", "shortlist_count"])
        for tp in portfolio_analytics.time_series_trends:
            writer.writerow([tp.period, tp.session_count, tp.resume_count, tp.avg_score or "N/A", tp.shortlist_count])

        csv_bytes = buf.getvalue().encode("utf-8")
        filename = f"portfolio_analytics_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        file_path = str(self.export_dir / filename)
        with open(file_path, "wb") as f:
            f.write(csv_bytes)

        return csv_bytes, file_path

    def export_comparison_analytics_csv(
        self, comp_analytics: SessionComparisonAnalytics
    ) -> Tuple[bytes, str]:
        buf = io.StringIO()
        writer = csv.writer(buf, lineterminator="\n")

        writer.writerow(["# HIREMATCH MULTI-SESSION COMPARISON EXPORT"])
        writer.writerow(["# GENERATED AT", comp_analytics.generated_at])
        writer.writerow(["# COMPATIBLE SESSIONS", "YES" if comp_analytics.is_compatible else "NO"])
        for w in comp_analytics.compatibility_warnings:
            writer.writerow(["# WARNING", w])
        writer.writerow([])

        writer.writerow(["--- SUMMARY COMPARISON TABLE ---"])
        writer.writerow([
            "session_id",
            "session_name",
            "job_title",
            "total_enrolled",
            "scored_count",
            "mean_score",
            "median_score",
            "mean_coverage_pct",
            "shortlisted_count",
            "shortlist_conversion_rate_pct",
            "config_version",
        ])
        for row in comp_analytics.comparison_summary_table:
            writer.writerow([
                row.get("session_id"),
                row.get("session_name"),
                row.get("job_title"),
                row.get("total_enrolled"),
                row.get("scored_count"),
                row.get("mean_score"),
                row.get("median_score"),
                row.get("mean_coverage"),
                row.get("shortlisted_count"),
                row.get("shortlist_conversion_rate"),
                row.get("config_version"),
            ])
        writer.writerow([])

        writer.writerow(["--- ALIGNED CRITERIA SUPPORT RATES (%) ---"])
        sess_ids = comp_analytics.session_ids
        header = ["criterion_name"] + [comp_analytics.session_names.get(sid, sid[:8]) for sid in sess_ids]
        writer.writerow(header)

        for c_name, rate_dict in comp_analytics.aligned_criteria_support_rates.items():
            row_data = [c_name] + [rate_dict.get(sid, 0.0) for sid in sess_ids]
            writer.writerow(row_data)

        csv_bytes = buf.getvalue().encode("utf-8")
        filename = f"session_comparison_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        file_path = str(self.export_dir / filename)
        with open(file_path, "wb") as f:
            f.write(csv_bytes)

        return csv_bytes, file_path
