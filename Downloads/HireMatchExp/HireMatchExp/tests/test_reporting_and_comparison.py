
import io
from pathlib import Path
import tempfile
import time
import unittest

try:
    import pymupdf as fitz
except ImportError:
    import fitz

import openpyxl

from config import AppConfig
from core.models.candidate import Candidate
from core.models.enums import (
    CriterionCategory,
    CriterionType,
    FileType,
    MatchStatus,
    ProcessingStatus,
    ScoreCompleteness,
    ShortlistReviewStatus,
)
from core.models.evaluation import CandidateEvaluationRecord, CriterionEvaluationResult
from core.models.job import Job
from core.models.report_models import (
    ExportFormat,
    GeneratedReportRecord,
    ReportType,
)
from core.models.scoring_models import (
    CandidateScoringRecord,
    CriterionContribution,
    RequiredCriteriaSummary,
    ScoringConfiguration,
    ScoringCriterion,
    ShortlistRecord,
)
from core.models.session import ScreeningSession
from core.reporting.candidate_report import CandidateReportGenerator
from core.reporting.comparison_service import CandidateComparisonService
from core.reporting.export_service import ReportExportService
from core.reporting.formatters import (
    CSVReportFormatter,
    PDFReportFormatter,
    XLSXReportFormatter,
)
from core.reporting.insights_service import RecruitmentInsightsService
from core.reporting.session_report import SessionSummaryReportGenerator
from core.reporting.shortlist_report import ShortlistReportGenerator
from core.storage.database import Database
from core.storage.eval_repository import EvaluationRecordRepository
from core.storage.report_repository import ReportRecordRepository
from core.storage.repositories import (
    CandidateRepository,
    JobRepository,
    SessionRepository,
)
from core.storage.scoring_repository import ScoringRecordRepository
from core.storage.shortlist_repository import ShortlistRepository


class TestReportingAndComparisonSuite(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        db_path = Path(self.temp_dir.name) / "test_reporting.db"
        self.db = Database(db_path)
        self.db.initialize_schema()

        self.config = AppConfig(
            base_dir=Path(self.temp_dir.name),
            db_path=db_path,
            reports_dir=Path(self.temp_dir.name) / "reports",
        )
        self.config.ensure_directories()

        self.job_repo = JobRepository(self.db)
        self.sess_repo = SessionRepository(self.db)
        self.cand_repo = CandidateRepository(self.db)
        self.eval_repo = EvaluationRecordRepository(self.db)
        self.scoring_repo = ScoringRecordRepository(self.db)
        self.shortlist_repo = ShortlistRepository(self.db)
        self.report_repo = ReportRecordRepository(self.db)

        self.job = self.job_repo.create(
            Job(
                title="Lead Cloud Engineer",
                department="Platform Infrastructure",
                original_description="Lead Cloud Engineer with AWS, Python, Kubernetes and CI/CD.",
                parsed_requirements={"skills": ["AWS", "Python", "Kubernetes", "CI/CD"]},
            )
        )
        self.session = self.sess_repo.create(
            ScreeningSession(job_id=self.job.id, session_name="Cloud Platform Q4", max_resume_count=100)
        )

        self.cands = []
        for i, name in enumerate(["alice_smith.pdf", "bob_jones.pdf", "charlie_brown.pdf"]):
            c = self.cand_repo.create(
                Candidate(
                    original_filename=name,
                    file_type=FileType.PDF,
                    file_hash=f"hash_{i}".ljust(64, "0"),
                    processing_status=ProcessingStatus.COMPLETED,
                )
            )
            self.cand_repo.update_structured_data(c.id, {
                "contact": {"name": name.split(".")[0].replace("_", " ").title(), "email": f"{name.split('.')[0]}@test.com", "phone": "+1-555-0100"},
                "skills": [{"name": "Python"}, {"name": "AWS"}],
                "experience": [{"job_title": "Senior Cloud Engineer", "employer": "TechCorp"}],
                "education": [{"qualification": "Bachelor of Science in CS", "institution": "State University"}],
            })
            self.cand_repo.enroll_in_session(c.id, self.session.id, status=ProcessingStatus.COMPLETED)
            self.cands.append(c)

        cc_alice = [
            CriterionContribution(criterion_id="c1", criterion_name="Python", category=CriterionCategory.TECHNICAL_SKILL, criterion_type=CriterionType.REQUIRED, configured_weight=0.5, support_value=1.0, normalized_contribution=50.0, match_status=MatchStatus.SUPPORTED, is_eligible=True, notes="Verified senior Python experience"),
            CriterionContribution(criterion_id="c2", criterion_name="AWS", category=CriterionCategory.TECHNICAL_SKILL, criterion_type=CriterionType.REQUIRED, configured_weight=0.5, support_value=0.8, normalized_contribution=40.0, match_status=MatchStatus.SUPPORTED, is_eligible=True, notes="AWS Solutions Architect certified"),
        ]
        s_alice = CandidateScoringRecord(
            candidate_id=self.cands[0].id,
            session_id=self.session.id,
            evaluation_record_id="eval_1",
            scoring_config_id="cfg_1",
            scoring_config_version="v1.0",
            overall_score=90.0,
            evidence_coverage=100.0,
            completeness_status=ScoreCompleteness.COMPLETE,
            criterion_contributions=cc_alice,
            explanation="Strong candidate with verified credentials.",
        )
        self.scoring_repo.save(s_alice)
        self.shortlist_repo.save(ShortlistRecord(
            session_id=self.session.id,
            candidate_id=self.cands[0].id,
            review_status=ShortlistReviewStatus.SHORTLISTED,
            reviewer_notes="Exceptional candidate for leadership role.",
        ))

        cc_bob = [
            CriterionContribution(criterion_id="c1", criterion_name="Python", category=CriterionCategory.TECHNICAL_SKILL, criterion_type=CriterionType.REQUIRED, configured_weight=0.5, support_value=0.5, normalized_contribution=25.0, match_status=MatchStatus.PARTIALLY_SUPPORTED, is_eligible=True, notes="Basic scripts mentioned"),
            CriterionContribution(criterion_id="c2", criterion_name="AWS", category=CriterionCategory.TECHNICAL_SKILL, criterion_type=CriterionType.REQUIRED, configured_weight=0.5, support_value=0.8, normalized_contribution=40.0, match_status=MatchStatus.SUPPORTED, is_eligible=True, notes="AWS deployment history"),
        ]
        s_bob = CandidateScoringRecord(
            candidate_id=self.cands[1].id,
            session_id=self.session.id,
            evaluation_record_id="eval_2",
            scoring_config_id="cfg_1",
            scoring_config_version="v1.0",
            overall_score=65.0,
            evidence_coverage=100.0,
            completeness_status=ScoreCompleteness.COMPLETE,
            criterion_contributions=cc_bob,
            explanation="Good cloud experience, Python depth uncertain.",
        )
        self.scoring_repo.save(s_bob)
        self.shortlist_repo.save(ShortlistRecord(
            session_id=self.session.id,
            candidate_id=self.cands[1].id,
            review_status=ShortlistReviewStatus.IN_REVIEW,
            reviewer_notes="Schedule technical phone screen.",
        ))

        cc_charlie = [
            CriterionContribution(criterion_id="c1", criterion_name="Python", category=CriterionCategory.TECHNICAL_SKILL, criterion_type=CriterionType.REQUIRED, configured_weight=0.5, support_value=0.5, normalized_contribution=25.0, match_status=MatchStatus.PARTIALLY_SUPPORTED, is_eligible=True, notes="Mentioned Python in coursework"),
            CriterionContribution(criterion_id="c2", criterion_name="AWS", category=CriterionCategory.TECHNICAL_SKILL, criterion_type=CriterionType.REQUIRED, configured_weight=0.5, support_value=0.0, normalized_contribution=0.0, match_status=MatchStatus.NOT_FOUND, is_eligible=True, notes="AWS not mentioned"),
        ]
        s_charlie = CandidateScoringRecord(
            candidate_id=self.cands[2].id,
            session_id=self.session.id,
            evaluation_record_id="eval_3",
            scoring_config_id="cfg_1",
            scoring_config_version="v1.0",
            overall_score=25.0,
            evidence_coverage=50.0,
            completeness_status=ScoreCompleteness.COMPLETE,
            criterion_contributions=cc_charlie,
            explanation="Missing critical AWS requirements.",
        )
        self.scoring_repo.save(s_charlie)

        self.export_service = ReportExportService(self.db, self.config)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_candidate_report_generation_from_persisted_records(self):
        generator = CandidateReportGenerator(self.db)
        report = generator.generate(self.cands[0].id, self.session.id, include_pii=False)

        self.assertEqual(report.candidate_id, self.cands[0].id)
        self.assertEqual(report.candidate_reference, f"Cand-{self.cands[0].id[:8].upper()}")
        self.assertEqual(report.job_title, "Lead Cloud Engineer")
        self.assertEqual(report.overall_score, 90.0)
        self.assertEqual(report.evidence_coverage, 100.0)
        self.assertEqual(report.review_status, "SHORTLISTED")
        self.assertEqual(report.reviewer_notes, "Exceptional candidate for leadership role.")
        self.assertEqual(len(report.criterion_items), 2)
        self.assertEqual(report.criterion_items[0].criterion_name, "Python")
        self.assertEqual(report.criterion_items[0].normalized_points, 50.0)

    def test_candidate_report_pii_quarantine(self):
        generator = CandidateReportGenerator(self.db)

        masked_report = generator.generate(self.cands[0].id, self.session.id, include_pii=False)
        self.assertIn("PROTECTED", masked_report.contact_info["email"])
        self.assertIn("PROTECTED", masked_report.contact_info["phone"])

        unmasked_report = generator.generate(self.cands[0].id, self.session.id, include_pii=True)
        self.assertEqual(unmasked_report.contact_info["email"], "alice_smith@test.com")
        self.assertEqual(unmasked_report.contact_info["phone"], "+1-555-0100")

    def test_session_summary_report_and_coverage_limitations(self):
        generator = SessionSummaryReportGenerator(self.db)
        sess_report = generator.generate(self.session.id)

        self.assertEqual(sess_report.total_resumes_enrolled, 3)
        self.assertEqual(sess_report.completed_processing_count, 3)
        self.assertEqual(sess_report.failed_processing_count, 0)
        self.assertEqual(sess_report.completed_evaluation_count, 3)

        sd = sess_report.score_distribution
        self.assertEqual(sd["total_scored"], 3)
        self.assertEqual(sd["max"], 90.0)
        self.assertEqual(sd["min"], 25.0)
        self.assertEqual(sd["median"], 65.0)

        sl = sess_report.shortlist_summary
        self.assertEqual(sl[ShortlistReviewStatus.SHORTLISTED.value], 1)
        self.assertEqual(sl[ShortlistReviewStatus.IN_REVIEW.value], 1)

    def test_candidate_comparison_compatibility_validation(self):
        comp_service = CandidateComparisonService(self.db)

        cand_ids = [self.cands[0].id, self.cands[1].id]
        comp_data = comp_service.compare_candidates(cand_ids, self.session.id)

        self.assertTrue(comp_data.compatibility.is_compatible)
        self.assertEqual(len(comp_data.compatibility.compatibility_warnings), 0)
        self.assertEqual(len(comp_data.aligned_criteria_rows), 2)

        py_row = next(r for r in comp_data.aligned_criteria_rows if r.criterion_name == "Python")
        self.assertEqual(py_row.candidate_evaluations[self.cands[0].id]["match_status"], "SUPPORTED")
        self.assertEqual(py_row.candidate_evaluations[self.cands[1].id]["match_status"], "PARTIALLY_SUPPORTED")
        self.assertTrue(py_row.has_meaningful_difference)

        s_cand2_diff_ver = CandidateScoringRecord(
            candidate_id=self.cands[2].id,
            session_id=self.session.id,
            evaluation_record_id="eval_3_diff",
            scoring_config_id="cfg_2",
            scoring_config_version="v2.0-custom",
            overall_score=70.0,
            evidence_coverage=80.0,
            completeness_status=ScoreCompleteness.COMPLETE,
            criterion_contributions=[
                CriterionContribution(criterion_id="c_diff", criterion_name="Kubernetes", category=CriterionCategory.TECHNICAL_SKILL, criterion_type=CriterionType.REQUIRED, configured_weight=1.0, support_value=0.7, normalized_contribution=70.0, match_status=MatchStatus.SUPPORTED)
            ]
        )
        self.scoring_repo.save(s_cand2_diff_ver)

        incomp_comp = comp_service.compare_candidates([self.cands[0].id, self.cands[2].id], self.session.id)
        self.assertFalse(incomp_comp.compatibility.is_compatible)
        self.assertTrue(any("Version Discrepancy" in w for w in incomp_comp.compatibility.compatibility_warnings))

    def test_shortlist_report_human_in_the_loop_audit(self):
        sl_generator = ShortlistReportGenerator(self.db)
        sl_data = sl_generator.generate(self.session.id)

        self.assertEqual(sl_data.total_shortlisted_count, 2)
        top_cand = sl_data.items[0]
        self.assertEqual(top_cand.review_status, ShortlistReviewStatus.SHORTLISTED)
        self.assertIn("Exceptional candidate", top_cand.reviewer_notes)
        self.assertIn("Candidates are never shortlisted or rejected automatically", sl_data.responsible_ai_disclaimer)

    def test_pdf_export_generation_and_page_structure(self):
        pdf_bytes, file_path, rec = self.export_service.export_candidate_report(
            self.cands[0].id, self.session.id, export_format=ExportFormat.PDF
        )
        self.assertTrue(pdf_bytes.startswith(b"%PDF-"))
        self.assertTrue(Path(file_path).exists())

        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        self.assertGreaterEqual(len(doc), 1)
        first_page_text = doc[0].get_text()
        self.assertIn("Candidate Evaluation Dossier", first_page_text)
        self.assertIn("Score:", first_page_text)
        self.assertIn("Page 1 of", first_page_text)
        doc.close()

        sess_pdf_bytes, _, _ = self.export_service.export_session_report(
            self.session.id, export_format=ExportFormat.PDF
        )
        sess_doc = fitz.open(stream=sess_pdf_bytes, filetype="pdf")
        self.assertIn("Screening Session Summary", sess_doc[0].get_text())
        sess_doc.close()

    def test_csv_export_formatting_and_missing_values(self):
        csv_bytes, file_path, _ = self.export_service.export_candidate_report(
            self.cands[0].id, self.session.id, export_format=ExportFormat.CSV
        )
        csv_text = csv_bytes.decode("utf-8")
        self.assertIn("# Candidate Reference", csv_text)
        self.assertIn("criterion_id,criterion_name,category", csv_text)
        self.assertIn("Python", csv_text)
        self.assertIn("50.0", csv_text)

        sess_csv_bytes, _, _ = self.export_service.export_session_report(
            self.session.id, export_format=ExportFormat.CSV
        )
        sess_csv_text = sess_csv_bytes.decode("utf-8")
        self.assertIn("candidate_id,candidate_reference,original_filename", sess_csv_text)

    def test_xlsx_export_multi_tab_structure(self):
        xlsx_bytes, file_path, _ = self.export_service.export_candidate_report(
            self.cands[0].id, self.session.id, export_format=ExportFormat.XLSX
        )
        self.assertTrue(len(xlsx_bytes) > 0)

        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        self.assertIn("Executive Summary", wb.sheetnames)
        self.assertIn("Criterion Breakdown", wb.sheetnames)
        self.assertIn("Extracted Profile", wb.sheetnames)

        ws_sum = wb["Executive Summary"]
        self.assertIn("Candidate Evaluation Dossier", ws_sum["A1"].value)

    def test_report_metadata_persistence_and_history(self):
        _, _, record = self.export_service.export_session_report(
            self.session.id, export_format=ExportFormat.PDF
        )

        persisted = self.report_repo.get(record.id)
        self.assertIsNotNone(persisted)
        self.assertEqual(persisted.report_type, ReportType.SESSION_SUMMARY)
        self.assertEqual(persisted.export_format, ExportFormat.PDF)
        self.assertEqual(persisted.session_id, self.session.id)

        history = self.report_repo.list_by_session(self.session.id)
        self.assertGreaterEqual(len(history), 1)

    def test_recruitment_insights_statistical_aggregation(self):
        insights_service = RecruitmentInsightsService(self.db)
        insights = insights_service.get_session_insights(self.session.id)

        sc_stats = insights["score_statistics"]
        self.assertEqual(sc_stats["total_scored"], 3)
        self.assertIn("percentiles", sc_stats)

        strengths = insights["top_candidate_strengths"]
        self.assertGreaterEqual(len(strengths), 1)
        self.assertEqual(strengths[0]["criterion_name"], "Python")

        gaps = insights["top_pool_skill_gaps"]
        self.assertGreaterEqual(len(gaps), 1)
        self.assertEqual(gaps[0]["criterion_name"], "AWS")

    def test_100_resume_session_report_performance(self):
        perf_session = self.sess_repo.create(
            ScreeningSession(job_id=self.job.id, session_name="100 Resumes Perf Report", max_resume_count=100)
        )

        scoring_batch = []
        cand_ids = []
        for i in range(100):
            c = self.cand_repo.create(
                Candidate(
                    original_filename=f"perf_candidate_{i:03d}.pdf",
                    file_type=FileType.PDF,
                    file_hash=f"perf_unique_hash_{i:04d}".ljust(64, "x"),
                    processing_status=ProcessingStatus.COMPLETED,
                )
            )
            self.cand_repo.enroll_in_session(c.id, perf_session.id, status=ProcessingStatus.COMPLETED)
            cand_ids.append(c.id)

            s_rec = CandidateScoringRecord(
                candidate_id=c.id,
                session_id=perf_session.id,
                evaluation_record_id=f"eval_p_{i}",
                scoring_config_id="cfg_perf",
                scoring_config_version="v1.0",
                overall_score=(i % 100) * 1.0,
                evidence_coverage=75.0,
                completeness_status=ScoreCompleteness.COMPLETE,
                criterion_contributions=[
                    CriterionContribution(criterion_id="c1", criterion_name="Python", category=CriterionCategory.TECHNICAL_SKILL, criterion_type=CriterionType.REQUIRED, configured_weight=0.6, support_value=(i % 10) / 10.0, normalized_contribution=((i % 10) / 10.0) * 60.0, match_status=MatchStatus.SUPPORTED if (i % 10) > 0 else MatchStatus.NOT_FOUND),
                    CriterionContribution(criterion_id="c2", criterion_name="AWS", category=CriterionCategory.TECHNICAL_SKILL, criterion_type=CriterionType.REQUIRED, configured_weight=0.4, support_value=((i + 3) % 10) / 10.0, normalized_contribution=(((i + 3) % 10) / 10.0) * 40.0, match_status=MatchStatus.PARTIALLY_SUPPORTED),
                ]
            )
            scoring_batch.append(s_rec)

        self.scoring_repo.save_batch(scoring_batch)

        start_time = time.time()
        sess_rep_data = self.export_service.session_generator.generate(perf_session.id)
        duration_summary = time.time() - start_time

        self.assertEqual(sess_rep_data.total_resumes_enrolled, 100)
        self.assertEqual(sess_rep_data.completed_evaluation_count, 100)
        self.assertLess(duration_summary, 2.0, f"100-resume summary took {duration_summary:.3f}s, expected < 2.0s")

        start_comp = time.time()
        comp_data = self.export_service.comparison_service.compare_candidates(cand_ids[:5], perf_session.id)
        duration_comp = time.time() - start_comp

        self.assertEqual(len(comp_data.compared_candidate_ids), 5)
        self.assertLess(duration_comp, 1.0, f"5-candidate comparison took {duration_comp:.3f}s, expected < 1.0s")


if __name__ == "__main__":
    unittest.main()
