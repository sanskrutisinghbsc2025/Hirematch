
from datetime import datetime, timezone
import io
from pathlib import Path
import tempfile
import time
from typing import Any, Dict, List
import unittest

import numpy as np
import openpyxl

from config import AppConfig
from core.analytics.models import (
    AnalyticsFilter,
    GapStatus,
    InsightCategory,
    InsightSeverity,
    PortfolioAnalytics,
    SessionAnalytics,
    SessionComparisonAnalytics,
)
from core.models.candidate import Candidate
from core.models.enums import (
    CriterionCategory,
    CriterionType,
    FileType,
    MatchStatus,
    ProcessingStatus,
    ScoreCompleteness,
    ShortlistReviewStatus,
)
from core.models.evaluation import CandidateEvaluationRecord, CriterionEvaluationResult
from core.models.job import Job
from core.models.report_models import ExportFormat
from core.models.scoring_models import (
    CandidateScoringRecord,
    CriterionContribution,
    ShortlistRecord,
)
from core.models.session import ScreeningSession
from core.services.session_orchestrator import ScreeningOrchestrator
from core.storage.database import Database
from core.storage.eval_repository import EvaluationRecordRepository
from core.storage.repositories import (
    CandidateRepository,
    CriterionRepository,
    JobRepository,
    SessionRepository,
)
from core.storage.scoring_repository import ScoringRecordRepository
from core.storage.shortlist_repository import ShortlistRepository


class TestBusinessAnalyticsSuite(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp_dir.name) / "test_analytics.db"
        self.db = Database(str(self.db_path))
        self.db.initialize_schema()

        self.config = AppConfig()
        self.config.reports_dir = Path(self.temp_dir.name) / "reports"
        self.config.ensure_directories()

        self.orchestrator = ScreeningOrchestrator(self.db, self.config)
        self.job_repo = JobRepository(self.db)
        self.sess_repo = SessionRepository(self.db)
        self.cand_repo = CandidateRepository(self.db)
        self.scoring_repo = ScoringRecordRepository(self.db)
        self.eval_repo = EvaluationRecordRepository(self.db)
        self.shortlist_repo = ShortlistRepository(self.db)

        self.job = self.job_repo.create(Job(
            title="Senior Distributed Systems Engineer",
            department="Infrastructure",
            original_description="Leading distributed cloud infrastructure using Python, Kubernetes, and AWS.",
            parsed_requirements={"requirements": ["Python", "Kubernetes", "AWS", "Distributed Systems"]},
        ))

        self.session_1 = self.sess_repo.create(ScreeningSession(
            job_id=self.job.id,
            session_name="Session 1 - Infrastructure Cohort",
            max_resume_count=100,
            created_at="2026-10-01T10:00:00Z",
            completed_at="2026-10-01T10:05:00Z",
        ))

        self.session_2 = self.sess_repo.create(ScreeningSession(
            job_id=self.job.id,
            session_name="Session 2 - Cloud Platform Cohort",
            max_resume_count=100,
            created_at="2026-10-02T14:00:00Z",
            completed_at="2026-10-02T14:08:00Z",
        ))

        self.cands_s1 = []
        for i, name in enumerate(["alice_smith", "bob_jones", "charlie_brown", "diana_prince"]):
            c = self.cand_repo.create(Candidate(
                original_filename=f"{name}.pdf",
                file_type=FileType.PDF,
                file_hash=f"hash_s1_{i:04d}".ljust(64, "x"),
                processing_status=ProcessingStatus.COMPLETED,
            ))
            skills = [{"name": "Python"}, {"name": "K8s"}] if i < 2 else [{"name": "Python"}]
            self.cand_repo.update_structured_data(c.id, {
                "contact": {"name": name.replace("_", " ").title()},
                "skills": skills,
            })
            self.cand_repo.enroll_in_session(c.id, self.session_1.id, status=ProcessingStatus.COMPLETED)
            self.cands_s1.append(c)

        scores_data = [
            (self.cands_s1[0].id, 90.0, 1.0, 1.0, ShortlistReviewStatus.SHORTLISTED),
            (self.cands_s1[1].id, 85.0, 1.0, 0.7, ShortlistReviewStatus.IN_REVIEW),
            (self.cands_s1[2].id, 45.0, 0.5, 0.0, ShortlistReviewStatus.NEEDS_REVIEW),
        ]
        for cid, ov_sc, py_sup, k8s_sup, sl_stat in scores_data:
            cc = [
                CriterionContribution(
                    criterion_id="c_py",
                    criterion_name="Python",
                    category=CriterionCategory.TECHNICAL_SKILL,
                    criterion_type=CriterionType.REQUIRED,
                    configured_weight=0.5,
                    support_value=py_sup,
                    normalized_contribution=py_sup * 50.0,
                    match_status=MatchStatus.SUPPORTED if py_sup > 0.6 else MatchStatus.PARTIALLY_SUPPORTED,
                ),
                CriterionContribution(
                    criterion_id="c_k8s",
                    criterion_name="Kubernetes",
                    category=CriterionCategory.TECHNICAL_SKILL,
                    criterion_type=CriterionType.REQUIRED,
                    configured_weight=0.5,
                    support_value=k8s_sup,
                    normalized_contribution=k8s_sup * 50.0,
                    match_status=MatchStatus.SUPPORTED if k8s_sup > 0.6 else (MatchStatus.PARTIALLY_SUPPORTED if k8s_sup > 0 else MatchStatus.NOT_FOUND),
                ),
            ]
            self.scoring_repo.save(CandidateScoringRecord(
                candidate_id=cid,
                session_id=self.session_1.id,
                evaluation_record_id=f"eval_{cid}",
                scoring_config_id="cfg_1",
                scoring_config_version="v1.0",
                overall_score=ov_sc,
                evidence_coverage=100.0 if k8s_sup > 0 else 50.0,
                completeness_status=ScoreCompleteness.COMPLETE,
                criterion_contributions=cc,
            ))
            self.shortlist_repo.save(ShortlistRecord(
                session_id=self.session_1.id,
                candidate_id=cid,
                review_status=sl_stat,
                reviewer_notes="Standard evaluation.",
            ))

            self.eval_repo.save_record(CandidateEvaluationRecord(
                candidate_id=cid,
                session_id=self.session_1.id,
                config_id="cfg_1",
                config_version="v1.0",
                overall_score=ov_sc,
                evidence_coverage=100.0 if k8s_sup > 0 else 50.0,
                evaluation_status="CALCULATED",
                criterion_results=[
                    CriterionEvaluationResult(
                        criterion_id="c_py",
                        criterion_name="Python",
                        criterion_type=CriterionType.REQUIRED,
                        weight=0.5,
                        match_status=MatchStatus.SUPPORTED if py_sup >= 1.0 else MatchStatus.PARTIALLY_SUPPORTED,
                        score=py_sup,
                        semantic_score=py_sup,
                    ),
                    CriterionEvaluationResult(
                        criterion_id="c_k8s",
                        criterion_name="Kubernetes",
                        criterion_type=CriterionType.REQUIRED,
                        weight=0.5,
                        match_status=MatchStatus.SUPPORTED if k8s_sup >= 1.0 else (MatchStatus.PARTIALLY_SUPPORTED if k8s_sup > 0 else MatchStatus.NOT_FOUND),
                        score=0.2 if k8s_sup > 0 else 0.0,
                        semantic_score=k8s_sup,
                    ),
                ],
            ))

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_portfolio_analytics_kpis_and_aggregations(self):
        portfolio = self.orchestrator.get_portfolio_analytics()

        self.assertGreaterEqual(portfolio.total_sessions, 2)
        self.assertEqual(portfolio.total_resumes_enrolled, 4)
        self.assertEqual(portfolio.completed_processing_count, 4)
        self.assertEqual(portfolio.failed_processing_count, 0)
        self.assertIsNotNone(portfolio.overall_mean_score)
        self.assertAlmostEqual(portfolio.overall_mean_score, 73.33, places=1)
        self.assertEqual(portfolio.overall_median_score, 85.0)
        self.assertIn("q1", portfolio.score_quartiles)
        self.assertIn("q3", portfolio.score_quartiles)

    def test_time_series_trend_calculations(self):
        portfolio = self.orchestrator.get_portfolio_analytics()

        self.assertGreaterEqual(len(portfolio.time_series_trends), 1)
        tp = portfolio.time_series_trends[0]
        self.assertTrue(len(tp.period) >= 7)
        self.assertGreaterEqual(tp.session_count, 1)

    def test_session_score_statistics_and_distributions(self):
        sess_analytics = self.orchestrator.get_session_analytics_profile(self.session_1.id)

        self.assertEqual(sess_analytics.session_id, self.session_1.id)
        self.assertEqual(sess_analytics.total_enrolled, 4)
        self.assertEqual(sess_analytics.total_scored, 3)
        self.assertEqual(sess_analytics.unscorable_count, 1)
        self.assertAlmostEqual(sess_analytics.score_mean, 73.33, places=1)
        self.assertEqual(sess_analytics.score_median, 85.0)
        self.assertGreater(sess_analytics.score_std_dev, 0.0)
        self.assertGreater(sess_analytics.score_variance, 0.0)
        self.assertIn("80-100", sess_analytics.score_buckets)
        self.assertEqual(sess_analytics.score_buckets["80-100"], 2)
        self.assertEqual(sess_analytics.score_buckets["40-59"], 1)

    def test_skill_coverage_normalization_and_accuracy(self):
        skill_items = self.orchestrator.get_skill_coverage_analytics(self.session_1.id)

        self.assertGreaterEqual(len(skill_items), 2)
        k8s_item = next((s for s in skill_items if s.canonical_name == "Kubernetes"), None)
        self.assertIsNotNone(k8s_item)
        self.assertEqual(k8s_item.canonical_name, "Kubernetes")
        self.assertEqual(k8s_item.candidate_evidence_count, 2)
        self.assertEqual(k8s_item.total_candidates, 4)
        self.assertEqual(k8s_item.pool_coverage_percentage, 50.0)
        self.assertEqual(k8s_item.gap_status, GapStatus.MODERATE_COVERAGE)

    def test_criterion_variance_and_divergence_metrics(self):
        crit_analytics = self.orchestrator.get_criterion_analytics(self.session_1.id)

        self.assertEqual(len(crit_analytics), 2)
        py_crit = next(c for c in crit_analytics if c.criterion_name == "Python")
        k8s_crit = next(c for c in crit_analytics if c.criterion_name == "Kubernetes")

        self.assertGreater(py_crit.variance, 0.0)
        self.assertGreaterEqual(k8s_crit.semantic_boost_count, 1)

    def test_traceable_insights_threshold_triggering(self):
        sess_analytics = self.orchestrator.get_session_analytics_profile(self.session_1.id)
        insights = sess_analytics.insights

        self.assertTrue(any(i.category in (InsightCategory.SEMANTIC_DISCOVERY, InsightCategory.ACUTE_SKILL_GAP, InsightCategory.LOW_SCORE_VARIATION) for i in insights))
        for ins in insights:
            self.assertTrue(len(ins.calculation_basis) > 0, "Insight must have an explicit calculation basis.")
            self.assertTrue(len(ins.suggested_recruiter_action) > 0, "Insight must have a recruiter action.")

    def test_session_comparison_and_compatibility(self):
        comp = self.orchestrator.compare_sessions_analytics([self.session_1.id, self.session_2.id])
        self.assertTrue(comp.is_compatible)
        self.assertEqual(len(comp.comparison_summary_table), 2)
        self.assertIn("Python", comp.aligned_criteria_support_rates)

        job_diff = self.job_repo.create(Job(
            title="Product Manager",
            department="Product",
            original_description="Leading product strategy.",
            parsed_requirements={"requirements": ["Roadmapping", "Agile"]},
        ))
        sess_diff = self.sess_repo.create(ScreeningSession(
            job_id=job_diff.id,
            session_name="Session 3 - PM Cohort",
        ))

        comp_incomp = self.orchestrator.compare_sessions_analytics([self.session_1.id, sess_diff.id])
        self.assertFalse(comp_incomp.is_compatible)
        self.assertTrue(any("Cross-Job Warning" in w for w in comp_incomp.compatibility_warnings))

    def test_analytics_filters_application(self):
        filtered_sess = self.orchestrator.get_session_analytics_profile(
            self.session_1.id, filters=AnalyticsFilter(min_score=80.0)
        )
        self.assertEqual(filtered_sess.total_scored, 2)
        self.assertAlmostEqual(filtered_sess.score_mean, 87.5, places=1)

        sl_filtered = self.orchestrator.get_session_analytics_profile(
            self.session_1.id, filters=AnalyticsFilter(shortlist_status="SHORTLISTED")
        )
        self.assertEqual(sl_filtered.total_scored, 1)
        self.assertEqual(sl_filtered.score_mean, 90.0)

    def test_incomplete_and_missing_data_resilience(self):
        empty_sess = self.sess_repo.create(ScreeningSession(
            job_id=self.job.id,
            session_name="Empty Session",
        ))
        empty_analytics = self.orchestrator.get_session_analytics_profile(empty_sess.id)

        self.assertEqual(empty_analytics.total_enrolled, 0)
        self.assertEqual(empty_analytics.total_scored, 0)
        self.assertIsNone(empty_analytics.score_mean)
        self.assertIsNone(empty_analytics.score_median)
        self.assertEqual(empty_analytics.score_std_dev, 0.0)
        self.assertEqual(len(empty_analytics.insights), 0)

    def test_csv_and_xlsx_analytics_export_structure(self):
        csv_bytes, csv_path = self.orchestrator.export_session_analytics(
            self.session_1.id, export_format=ExportFormat.CSV
        )
        csv_text = csv_bytes.decode("utf-8")
        self.assertIn("# METRIC DEFINITIONS", csv_text)
        self.assertIn("--- CRITERION PERFORMANCE & VARIANCE ---", csv_text)
        self.assertIn("Python", csv_text)
        self.assertTrue(Path(csv_path).exists())

        xlsx_bytes, xlsx_path = self.orchestrator.export_session_analytics(
            self.session_1.id, export_format=ExportFormat.XLSX
        )
        self.assertTrue(len(xlsx_bytes) > 0)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        self.assertIn("Executive Overview", wb.sheetnames)
        self.assertIn("Criterion Variance & AI", wb.sheetnames)
        self.assertIn("Skill Coverage & Gaps", wb.sheetnames)
        self.assertIn("Screening Insights", wb.sheetnames)
        self.assertIn("Metric Definitions", wb.sheetnames)
        self.assertTrue(Path(xlsx_path).exists())

        p_csv_bytes, p_csv_path = self.orchestrator.export_portfolio_analytics()
        p_csv_text = p_csv_bytes.decode("utf-8")
        self.assertIn("# HIREMATCH PORTFOLIO ANALYTICS EXPORT", p_csv_text)
        self.assertIn("Total Resumes Enrolled", p_csv_text)

    def test_100_resume_session_analytics_performance(self):
        perf_session = self.sess_repo.create(ScreeningSession(
            job_id=self.job.id,
            session_name="100 Resumes Analytics Performance Test",
            max_resume_count=100,
        ))

        scoring_batch = []
        for i in range(100):
            c = self.cand_repo.create(Candidate(
                original_filename=f"perf_cand_{i:03d}.pdf",
                file_type=FileType.PDF,
                file_hash=f"perf_analytics_hash_{i:04d}".ljust(64, "z"),
                processing_status=ProcessingStatus.COMPLETED,
            ))
            self.cand_repo.enroll_in_session(c.id, perf_session.id, status=ProcessingStatus.COMPLETED)

            scoring_batch.append(CandidateScoringRecord(
                candidate_id=c.id,
                session_id=perf_session.id,
                evaluation_record_id=f"eval_perf_{i}",
                scoring_config_id="cfg_perf",
                scoring_config_version="v1.0",
                overall_score=float((i * 7) % 100),
                evidence_coverage=75.0,
                completeness_status=ScoreCompleteness.COMPLETE,
                criterion_contributions=[
                    CriterionContribution(
                        criterion_id="c_py",
                        criterion_name="Python",
                        category=CriterionCategory.TECHNICAL_SKILL,
                        criterion_type=CriterionType.REQUIRED,
                        configured_weight=0.6,
                        support_value=(i % 10) / 10.0,
                        normalized_contribution=((i % 10) / 10.0) * 60.0,
                        match_status=MatchStatus.SUPPORTED if (i % 10) > 4 else MatchStatus.NOT_FOUND,
                    ),
                    CriterionContribution(
                        criterion_id="c_k8s",
                        criterion_name="Kubernetes",
                        category=CriterionCategory.TECHNICAL_SKILL,
                        criterion_type=CriterionType.REQUIRED,
                        configured_weight=0.4,
                        support_value=((i + 3) % 10) / 10.0,
                        normalized_contribution=(((i + 3) % 10) / 10.0) * 40.0,
                        match_status=MatchStatus.PARTIALLY_SUPPORTED,
                    ),
                ],
            ))

        self.scoring_repo.save_batch(scoring_batch)

        start_time = time.time()
        s_analytics = self.orchestrator.get_session_analytics_profile(perf_session.id)
        duration = time.time() - start_time

        self.assertEqual(s_analytics.total_enrolled, 100)
        self.assertEqual(s_analytics.total_scored, 100)
        self.assertLess(duration, 1.0, f"100-resume analytics took {duration:.3f}s, expected < 1.0s")


if __name__ == "__main__":
    unittest.main()
